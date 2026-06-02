"""Generate an Excel report of the SO/PO Playwright E2E run (2026-06-02)."""
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = Workbook()

# ── Styles ───────────────────────────────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
PASS_FILL   = PatternFill("solid", fgColor="DCFCE7")
PASS_FONT   = Font(bold=True, color="166534")
FAIL_FILL   = PatternFill("solid", fgColor="FEE2E2")
FAIL_FONT   = Font(bold=True, color="991B1B")
TITLE_FONT  = Font(bold=True, size=14, color="111827")
thin = Side(style="thin", color="D1D5DB")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(vertical="top", wrap_text=True)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        cell.border = BORDER


def mark_status(cell):
    if cell.value == "PASS":
        cell.fill, cell.font = PASS_FILL, PASS_FONT
    elif cell.value == "FAIL":
        cell.fill, cell.font = FAIL_FILL, FAIL_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")


# ══════════════════════════════════════════════════════════════════════════════
# SHEET 1 — Endpoint Coverage
# ══════════════════════════════════════════════════════════════════════════════
ws = wb.active
ws.title = "Endpoint Coverage"

ws["A1"] = "Skarpine — Sales Order & Purchase Order Endpoint Test Coverage"
ws["A1"].font = TITLE_FONT
ws["A2"] = f"Run: 2026-06-02  ·  Playwright E2E  ·  backend http://localhost:3001/api/v1"
ws["A2"].font = Font(italic=True, color="6B7280")

headers = ["#", "Method", "Endpoint", "Process", "Covered by test", "Status", "Notes"]
hr = 4
for i, h in enumerate(headers, 1):
    ws.cell(row=hr, column=i, value=h)
style_header(ws, hr, len(headers))

endpoints = [
    ("POST", "/auth/login", "Auth", "API setup (apiLogin helper)", "PASS", ""),
    ("GET",  "/products", "Sales/Purchase", "SO & PO lifecycle setup", "PASS", ""),
    # Sales
    ("POST", "/sales/orders", "Sales Order", "SO full lifecycle", "PASS", "Was 500 — fixed TAX.iva bug"),
    ("GET",  "/sales/orders", "Sales Order", "SO list / detail setup", "PASS", ""),
    ("GET",  "/sales/orders/:id", "Sales Order", "SO detail page loads", "PASS", ""),
    ("POST", "/sales/orders/:id/confirm", "Sales Order", "SO full lifecycle", "PASS", "DRAFT → CONFIRMED, reserves stock"),
    ("POST", "/sales/orders/:id/invoice", "Sales Order", "SO full lifecycle", "PASS", "Factura + GL JE; IVA13%/IT3%"),
    ("POST", "/sales/orders/:id/pay", "Sales Order", "SO full lifecycle", "PASS", "AR payment, clears CxC 1103"),
    # Purchase
    ("GET",  "/purchase/suppliers", "Purchase Order", "PO lifecycle + suppliers UI", "PASS", ""),
    ("POST", "/purchase/suppliers", "Purchase Order", "PO lifecycle (if no supplier)", "PASS", "Creates test supplier on demand"),
    ("POST", "/purchase/orders", "Purchase Order", "PO full lifecycle", "PASS", "Was 500 — fixed TAX.iva + counter"),
    ("POST", "/purchase/orders/:id/confirm", "Purchase Order", "PO full lifecycle", "PASS", "DRAFT → CONFIRMED"),
    ("POST", "/purchase/orders/:id/receive", "Purchase Order", "PO full lifecycle", "PASS", "Adds stock + GL JE (inv/IVA/AP)"),
    ("POST", "/purchase/orders/:id/pay", "Purchase Order", "PO full lifecycle", "PASS", "AP payment, clears CxP 2101"),
    # Support endpoints used in setup
    ("GET",  "/warehouse/warehouses", "Warehouse", "PO lifecycle setup", "PASS", ""),
    ("GET",  "/warehouse/locations", "Warehouse", "PO receive setup", "PASS", ""),
]

r = hr + 1
for idx, (method, ep, proc, test, status, notes) in enumerate(endpoints, 1):
    ws.cell(row=r, column=1, value=idx)
    ws.cell(row=r, column=2, value=method)
    ws.cell(row=r, column=3, value=ep)
    ws.cell(row=r, column=4, value=proc)
    ws.cell(row=r, column=5, value=test)
    sc = ws.cell(row=r, column=6, value=status)
    mark_status(sc)
    ws.cell(row=r, column=7, value=notes)
    for c in range(1, 8):
        ws.cell(row=r, column=c).border = BORDER
        if c in (5, 7):
            ws.cell(row=r, column=c).alignment = WRAP
    r += 1

# Summary line
ws.cell(row=r + 1, column=2, value="TOTAL")
ws.cell(row=r + 1, column=2).font = Font(bold=True)
ws.cell(row=r + 1, column=3, value=f"{len(endpoints)} endpoints")
ws.cell(row=r + 1, column=6, value="16/16 PASS")
ws.cell(row=r + 1, column=6).font = PASS_FONT

widths = {"A": 5, "B": 8, "C": 30, "D": 15, "E": 28, "F": 9, "G": 38}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

# ══════════════════════════════════════════════════════════════════════════════
# SHEET 2 — Test Cases (17 Playwright tests)
# ══════════════════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("Test Cases")
ws2["A1"] = "Playwright Test Cases — 06-sales-orders + 07-purchase-orders"
ws2["A1"].font = TITLE_FONT
ws2["A2"] = "17 tests · 17 passed · ~1.1 min · chromium"
ws2["A2"].font = Font(italic=True, color="6B7280")

headers2 = ["#", "Spec file", "Test name", "Type", "Status"]
hr2 = 4
for i, h in enumerate(headers2, 1):
    ws2.cell(row=hr2, column=i, value=h)
style_header(ws2, hr2, len(headers2))

tests = [
    ("06-sales-orders", "list page loads with table and New Order button", "UI", "PASS"),
    ("06-sales-orders", "status filter dropdown is visible", "UI", "PASS"),
    ("06-sales-orders", "New Order shows the SO create form", "UI", "PASS"),
    ("06-sales-orders", "Create Sales Order button disabled until product selected", "UI", "PASS"),
    ("06-sales-orders", "Cancel on SO form returns to list", "UI", "PASS"),
    ("06-sales-orders", "SO form has Customer, Warehouse, Discount, Notes fields", "UI", "PASS"),
    ("06-sales-orders", "Add line button adds a new product row", "UI", "PASS"),
    ("06-sales-orders", "SO lifecycle: DRAFT → CONFIRMED → Invoiced → Paid", "Lifecycle", "PASS"),
    ("06-sales-orders", "SO detail page loads for an existing order", "UI", "PASS"),
    ("07-purchase-orders", "list page loads with table and New PO button", "UI", "PASS"),
    ("07-purchase-orders", "table columns are visible", "UI", "PASS"),
    ("07-purchase-orders", "New PO button shows the create form", "UI", "PASS"),
    ("07-purchase-orders", "Create PO button disabled until supplier+warehouse+product", "UI", "PASS"),
    ("07-purchase-orders", "Cancel on PO form returns to list", "UI", "PASS"),
    ("07-purchase-orders", "PO form has Supplier, Warehouse, Expected Date fields", "UI", "PASS"),
    ("07-purchase-orders", "PO lifecycle: DRAFT → CONFIRMED → RECEIVED → Paid", "Lifecycle", "PASS"),
    ("07-purchase-orders", "suppliers page loads", "UI", "PASS"),
]

r = hr2 + 1
for idx, (spec, name, typ, status) in enumerate(tests, 1):
    ws2.cell(row=r, column=1, value=idx)
    ws2.cell(row=r, column=2, value=spec)
    ws2.cell(row=r, column=3, value=name)
    ws2.cell(row=r, column=4, value=typ)
    sc = ws2.cell(row=r, column=5, value=status)
    mark_status(sc)
    for c in range(1, 6):
        ws2.cell(row=r, column=c).border = BORDER
        if c == 3:
            ws2.cell(row=r, column=c).alignment = WRAP
    r += 1

ws2.cell(row=r + 1, column=2, value="TOTAL")
ws2.cell(row=r + 1, column=2).font = Font(bold=True)
ws2.cell(row=r + 1, column=3, value="17 tests")
ws2.cell(row=r + 1, column=5, value="17/17 PASS")
ws2.cell(row=r + 1, column=5).font = PASS_FONT

w2 = {"A": 5, "B": 20, "C": 55, "D": 12, "E": 9}
for col, w in w2.items():
    ws2.column_dimensions[col].width = w

# ══════════════════════════════════════════════════════════════════════════════
# SHEET 3 — Bugs Found & Fixed
# ══════════════════════════════════════════════════════════════════════════════
ws3 = wb.create_sheet("Bugs Fixed")
ws3["A1"] = "Bugs Uncovered by the SO/PO E2E Tests"
ws3["A1"].font = TITLE_FONT

headers3 = ["#", "Bug", "Severity", "Root cause", "Fix", "Status"]
hr3 = 3
for i, h in enumerate(headers3, 1):
    ws3.cell(row=hr3, column=i, value=h)
style_header(ws3, hr3, len(headers3))

bugs = [
    ("1", "TAX.iva is not a function — every SO/PO create returned HTTP 500",
     "Critical",
     "Refactor changed resolveTax() to return vat()/secondary(), but sales.service.ts:33 and purchase.routes.ts still called old TAX.iva()/TAX.it().",
     "config/tax.ts: resolveTax() now also returns iva:vat, it:secondary, IVA_RATE, IT_RATE (backward-compat). 17/17 tax unit tests pass.",
     "FIXED"),
    ("2", "order_counters desync — order_number unique constraint violation",
     "High",
     "order_counters started at 1 but seed data already used those SO/PO numbers. factura_counters self-heals with GREATEST(); order_counters did not.",
     "shared/utils/orderCounter.ts: nextCounter() uses GREATEST(last_number+1, currentMax+1); maxDocSuffix() reads current max from sales_orders/purchase_orders.",
     "FIXED"),
]

r = hr3 + 1
for (num, bug, sev, cause, fix, status) in bugs:
    ws3.cell(row=r, column=1, value=num)
    ws3.cell(row=r, column=2, value=bug)
    ws3.cell(row=r, column=3, value=sev)
    ws3.cell(row=r, column=4, value=cause)
    ws3.cell(row=r, column=5, value=fix)
    sc = ws3.cell(row=r, column=6, value=status)
    sc.fill, sc.font = PASS_FILL, PASS_FONT
    sc.alignment = Alignment(horizontal="center", vertical="center")
    for c in range(1, 7):
        ws3.cell(row=r, column=c).border = BORDER
        if c in (2, 4, 5):
            ws3.cell(row=r, column=c).alignment = WRAP
    ws3.row_dimensions[r].height = 75
    r += 1

w3 = {"A": 5, "B": 32, "C": 10, "D": 45, "E": 45, "F": 9}
for col, w in w3.items():
    ws3.column_dimensions[col].width = w

# ── Save ─────────────────────────────────────────────────────────────────────
out = r"C:\Users\Nieuw\Desktop\ClaUDE\wEBINAR\skarpine\docs\SO-PO-Test-Report.xlsx"
wb.save(out)
print("Saved:", out)
